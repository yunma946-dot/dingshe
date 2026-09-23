from __future__ import annotations

import json
import re
import subprocess
import sys
from datetime import date, datetime
from pathlib import Path


ROOT = Path(__file__).resolve().parent
DEFAULT_WORKBOOK = ROOT / "顶奢_整站资料批量编辑.xlsx"
SITE_CONTENT_PATH = ROOT / "site_content.json"
PROFILE_SHEET = "资料批量编辑"
HOME_SHEET = "首页设置"
CITY_SHEET = "城市首页"
PAGE_SHEET = "固定页面"
FOOTER_SHEET = "页脚设置"
DOMAIN = "https://hqq2.com"
IMAGE_EXTENSIONS = ("jpg", "jpeg", "png", "webp", "avif")
AUTO_PROFILE_TEXT = {
    "SEO描述": ("浏览", "的独立资料、图片说明、内容介绍与相关推荐。"),
    "正文1": ("以上是", "的照片、资料、无需注册登录即可浏览。"),
    "图片1 ALT": ("", "主图"),
    "图片1说明": ("", "资料主图"),
    "图片2 ALT": ("", "细节图"),
    "图片2说明": ("", "资料细节"),
    "图片3 ALT": ("", "展示图"),
    "图片3说明": ("", "资料展示"),
    "视频标题": ("", "资料视频"),
    "视频描述": ("", "相关资料视频"),
}

PROFILE_HEADERS = [
    "城市", "city_slug", "编号", "slug", "页面名称", "SEO标题", "SEO描述",
    "卡片/顶部简介", "风格/标签", "详情说明", "正文1", "正文2", "正文3",
    "图片文件夹", "图片1", "图片2", "图片3", "视频文件", "处理图片", "页面URL",
    "发布状态", "允许收录", "首页精选", "显示顺序", "首次发布日期", "最后更新日期",
    "图片1 ALT", "图片1说明", "图片2 ALT", "图片2说明", "图片3 ALT", "图片3说明",
    "视频标题", "视频描述", "视频封面", "视频上传日期", "视频时长", "内容审核状态",
]

CITY_FIELDS = [
    ("城市名称", "name"), ("城市代码", "slug"), ("导航短句", "line"),
    ("SEO标题", "seo_title"), ("SEO描述", "seo_description"),
    ("首屏英文", "hero_kicker"), ("首屏标题", "hero_title"), ("首屏简介", "intro"),
    ("数量标签", "count_label"), ("专题英文", "editorial_kicker"),
    ("专题标题", "editorial_title"), ("正文1", "paragraph_1"), ("正文2", "paragraph_2"),
    ("亮点1标题", "highlight_1_title"), ("亮点1说明", "highlight_1_text"),
    ("亮点2标题", "highlight_2_title"), ("亮点2说明", "highlight_2_text"),
    ("亮点3标题", "highlight_3_title"), ("亮点3说明", "highlight_3_text"),
    ("精选英文", "featured_kicker"), ("精选标题", "featured_title"),
    ("精选说明", "featured_description"), ("全部英文", "all_kicker"),
    ("全部标题", "all_title"), ("全部说明", "all_description"),
    ("FAQ英文", "faq_kicker"), ("FAQ标题", "faq_title"),
    ("FAQ说明", "faq_description"), ("问题1", "faq_1_question"),
    ("回答1", "faq_1_answer"), ("问题2", "faq_2_question"),
    ("回答2", "faq_2_answer"), ("问题3", "faq_3_question"),
    ("回答3", "faq_3_answer"), ("更多城市英文", "more_kicker"),
    ("更多城市标题", "more_title"),
    ("发布状态", "publish_status"), ("允许收录", "allow_index"),
    ("显示顺序", "display_order"), ("最后更新日期", "last_updated"),
]

PUBLISH_STATUSES = {"草稿", "已发布", "已下线"}
REVIEW_STATUSES = {"待完善", "待审核", "已完成"}
YES_NO = {"是", "否"}
TEMPLATE_MARKERS = ("当前为首批资料模板", "待更新", "示例内容", "预留三张照片", "预留三张图片")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def date_text(value: object) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if not text:
        return ""
    match = re.fullmatch(r"(\d{4})[-/.年](\d{1,2})[-/.月](\d{1,2})(?:日)?(?:\s+00:00:00)?", text)
    if not match:
        raise ValueError(f"日期格式不正确：{text}，请使用 YYYY-MM-DD")
    return date(int(match.group(1)), int(match.group(2)), int(match.group(3))).isoformat()


def positive_int(value: object, label: str, row: int) -> int:
    text = cell_text(value)
    try:
        number = int(float(text))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"第 {row} 行“{label}”必须是正整数。") from exc
    if number < 1:
        raise ValueError(f"第 {row} 行“{label}”必须大于 0。")
    return number


def find_header_row(sheet, required: set[str]) -> int:
    for row_number in range(1, min(sheet.max_row, 12) + 1):
        values = {
            cell_text(sheet.cell(row_number, column).value)
            for column in range(1, sheet.max_column + 1)
        }
        if required.issubset(values):
            return row_number
    raise ValueError(f"工作表“{sheet.title}”没有找到标题行，请不要修改表格结构。")


def column_map(sheet, header_row: int) -> dict[str, int]:
    return {
        cell_text(sheet.cell(header_row, column).value): column
        for column in range(1, sheet.max_column + 1)
        if cell_text(sheet.cell(header_row, column).value)
    }


def require_sheets(workbook) -> None:
    expected = [PROFILE_SHEET, HOME_SHEET, CITY_SHEET, PAGE_SHEET, FOOTER_SHEET]
    missing = [name for name in expected if name not in workbook.sheetnames]
    if missing:
        raise ValueError("工作表缺少分页：" + "、".join(missing))


def sync_profile_auto_formulas(workbook) -> bool:
    from openpyxl.utils import get_column_letter

    sheet = workbook[PROFILE_SHEET]
    header_row = find_header_row(sheet, {"页面名称", *AUTO_PROFILE_TEXT})
    columns = column_map(sheet, header_row)
    name_column = get_column_letter(columns["页面名称"])
    changed = False
    for row in range(header_row + 1, sheet.max_row + 1):
        if not cell_text(sheet.cell(row, columns["页面名称"]).value):
            continue
        name_reference = f"${name_column}{row}"
        for field, (prefix, suffix) in AUTO_PROFILE_TEXT.items():
            text_parts = [f'"{prefix}"'] if prefix else []
            text_parts.extend([name_reference, f'"{suffix}"'])
            formula = f'=IF({name_reference}="","",{"&".join(text_parts)})'
            cell = sheet.cell(row, columns[field])
            if cell.value != formula:
                cell.value = formula
                changed = True
    if changed:
        workbook.calculation.calcMode = "auto"
        workbook.calculation.fullCalcOnLoad = True
        workbook.calculation.forceFullCalc = True
    return changed


def read_settings_sheet(workbook, sheet_name: str, expected_fields: set[str]) -> dict[str, str]:
    sheet = workbook[sheet_name]
    header_row = find_header_row(sheet, {"字段代码", "内容"})
    columns = column_map(sheet, header_row)
    result: dict[str, str] = {}
    for row in range(header_row + 1, sheet.max_row + 1):
        field_code = cell_text(sheet.cell(row, columns["字段代码"]).value)
        if not field_code:
            continue
        if field_code in result:
            raise ValueError(f"工作表“{sheet_name}”字段代码重复：{field_code}")
        result[field_code] = cell_text(sheet.cell(row, columns["内容"]).value)
    missing = sorted(expected_fields - set(result))
    if missing:
        raise ValueError(f"工作表“{sheet_name}”缺少字段：" + "、".join(missing))
    return result


def read_cities(workbook, existing_cities: list[dict]) -> list[dict]:
    sheet = workbook[CITY_SHEET]
    required_headers = [label for label, _ in CITY_FIELDS]
    header_row = find_header_row(sheet, {"城市名称", "城市代码", "SEO标题"})
    columns = column_map(sheet, header_row)
    missing_headers = [header for header in required_headers if header not in columns]
    if missing_headers:
        raise ValueError("工作表“城市首页”缺少字段：" + "、".join(missing_headers))

    cities: list[dict] = []
    seen: set[str] = set()
    for row in range(header_row + 1, sheet.max_row + 1):
        name = cell_text(sheet.cell(row, columns["城市名称"]).value)
        slug = cell_text(sheet.cell(row, columns["城市代码"]).value).lower()
        if not name and not slug:
            continue
        if not name:
            raise ValueError(f"工作表“城市首页”第 {row} 行城市名称不能为空。")
        if not re.fullmatch(r"[a-z0-9-]+", slug):
            raise ValueError(f"工作表“城市首页”第 {row} 行城市代码不正确：{slug}")
        if slug in seen:
            raise ValueError(f"工作表“城市首页”城市代码重复：{slug}")
        seen.add(slug)
        city = {
            field_code: cell_text(sheet.cell(row, columns[label]).value)
            for label, field_code in CITY_FIELDS
        }
        city["last_updated"] = date_text(sheet.cell(row, columns["最后更新日期"]).value)
        city["display_order"] = positive_int(sheet.cell(row, columns["显示顺序"]).value, "显示顺序", row)
        if city["publish_status"] not in PUBLISH_STATUSES:
            raise ValueError(f"城市首页第 {row} 行发布状态必须是：草稿、已发布或已下线。")
        if city["allow_index"] not in YES_NO:
            raise ValueError(f"城市首页第 {row} 行允许收录必须是“是”或“否”。")
        if city["allow_index"] == "是" and city["publish_status"] != "已发布":
            raise ValueError(f"城市首页第 {row} 行只有“已发布”城市才能允许收录。")
        cities.append(city)

    expected_codes = {city["slug"] for city in existing_cities}
    actual_codes = {city["slug"] for city in cities}
    missing_codes = sorted(expected_codes - actual_codes)
    if missing_codes:
        raise ValueError("已存在城市的代码不能删除或修改：" + "、".join(missing_codes))
    cities.sort(key=lambda item: (item["display_order"], item["slug"]))
    return cities


def read_fixed_pages(workbook, existing_pages: dict[str, dict]) -> dict[str, dict]:
    sheet = workbook[PAGE_SHEET]
    header_row = find_header_row(sheet, {"page_code", "字段代码", "内容"})
    columns = column_map(sheet, header_row)
    pages: dict[str, dict] = {code: {} for code in existing_pages}
    seen: set[tuple[str, str]] = set()
    for row in range(header_row + 1, sheet.max_row + 1):
        page_code = cell_text(sheet.cell(row, columns["page_code"]).value)
        field_code = cell_text(sheet.cell(row, columns["字段代码"]).value)
        if not page_code and not field_code:
            continue
        if page_code not in pages:
            raise ValueError(f"工作表“固定页面”第 {row} 行 page_code 不正确：{page_code}")
        key = (page_code, field_code)
        if not field_code or key in seen:
            raise ValueError(f"工作表“固定页面”第 {row} 行字段代码为空或重复。")
        seen.add(key)
        pages[page_code][field_code] = cell_text(sheet.cell(row, columns["内容"]).value)

    for page_code, current in existing_pages.items():
        missing = sorted(set(current) - set(pages[page_code]))
        if missing:
            raise ValueError(f"固定页面 {page_code} 缺少字段：" + "、".join(missing))
    return pages


def read_profiles(workbook, cities: list[dict], brand_name: str) -> list[dict]:
    sheet = workbook[PROFILE_SHEET]
    header_row = find_header_row(sheet, {"city_slug", "slug", "页面名称"})
    columns = column_map(sheet, header_row)
    missing_headers = [header for header in PROFILE_HEADERS if header not in columns]
    if missing_headers:
        raise ValueError("工作表“资料批量编辑”缺少字段：" + "、".join(missing_headers))

    city_by_slug = {city["slug"]: city for city in cities}
    city_names = {slug: city["name"] for slug, city in city_by_slug.items()}
    city_order = {city["slug"]: city["display_order"] for city in cities}
    profiles: list[dict] = []
    seen_paths: set[str] = set()
    media_warnings: list[str] = []

    def value(row: int, header: str) -> str:
        return cell_text(sheet.cell(row, columns[header]).value)

    for row in range(header_row + 1, sheet.max_row + 1):
        city_slug = value(row, "city_slug").lower()
        slug = value(row, "slug").lower()
        city_cell = value(row, "城市")
        if not city_slug and not slug and not city_cell:
            continue
        if city_slug not in city_names:
            raise ValueError(f"资料表第 {row} 行 city_slug 不正确：{city_slug}")
        if not re.fullmatch(r"model-\d{3}", slug):
            raise ValueError(f"资料表第 {row} 行 slug 必须类似 model-001：{slug}")

        city = city_names[city_slug]
        if city_cell and city_cell != city:
            raise ValueError(f"资料表第 {row} 行城市名称与 city_slug 不匹配：{city_cell} / {city_slug}")
        path = f"{city_slug}/{slug}/"
        if path in seen_paths:
            raise ValueError(f"资料表第 {row} 行页面重复：{path}")
        seen_paths.add(path)

        name = value(row, "页面名称")
        if not name:
            raise ValueError(f"资料表第 {row} 行页面名称不能为空。")
        summary = value(row, "卡片/顶部简介") or f"{name}独立资料展示页。"
        seo_title = value(row, "SEO标题") or f"{name}｜图片与视频｜{brand_name}"
        seo_description = f"浏览{name}的独立资料、图片说明、内容介绍与相关推荐。"
        media_dir = f"photos/{city_slug}-{slug}"
        process_media = value(row, "处理图片") == "是"
        publish_status = value(row, "发布状态")
        allow_index = value(row, "允许收录")
        home_featured = value(row, "首页精选")
        content_review = value(row, "内容审核状态")
        if publish_status not in PUBLISH_STATUSES:
            raise ValueError(f"资料表第 {row} 行发布状态必须是：草稿、已发布或已下线。")
        if allow_index not in YES_NO or home_featured not in YES_NO:
            raise ValueError(f"资料表第 {row} 行允许收录、首页精选必须是“是”或“否”。")
        if content_review not in REVIEW_STATUSES:
            raise ValueError(f"资料表第 {row} 行内容审核状态必须是：待完善、待审核或已完成。")
        if allow_index == "是" and publish_status != "已发布":
            raise ValueError(f"资料表第 {row} 行只有“已发布”资料才能允许收录。")
        if home_featured == "是" and publish_status != "已发布":
            raise ValueError(f"资料表第 {row} 行只有“已发布”资料才能设为首页精选。")

        image_alts = {1: f"{name}主图", 2: f"{name}细节图", 3: f"{name}展示图"}
        image_captions = {1: f"{name}资料主图", 2: f"{name}资料细节", 3: f"{name}资料展示"}
        first_published = date_text(sheet.cell(row, columns["首次发布日期"]).value)
        last_updated = date_text(sheet.cell(row, columns["最后更新日期"]).value)
        video_upload_date = date_text(sheet.cell(row, columns["视频上传日期"]).value)

        profile = {
            "city": city,
            "city_slug": city_slug,
            "slug": slug,
            "name": name,
            "tag": value(row, "风格/标签") or "精选",
            "summary": summary,
            "intro": value(row, "详情说明"),
            "seo_title": seo_title,
            "seo_description": seo_description,
            "body_1": f"以上是{name}的照片、资料、无需注册登录即可浏览。",
            "body_2": value(row, "正文2"),
            "body_3": value(row, "正文3"),
            "process_media": process_media,
            "publish_status": publish_status,
            "allow_index": allow_index,
            "home_featured": home_featured,
            "display_order": positive_int(sheet.cell(row, columns["显示顺序"]).value, "显示顺序", row),
            "first_published": first_published,
            "last_updated": last_updated,
            "image_1_alt": image_alts[1],
            "image_1_caption": image_captions[1],
            "image_2_alt": image_alts[2],
            "image_2_caption": image_captions[2],
            "image_3_alt": image_alts[3],
            "image_3_caption": image_captions[3],
            "video_title": f"{name}资料视频",
            "video_description": f"{name}相关资料视频",
            "video_poster": value(row, "视频封面"),
            "video_upload_date": video_upload_date,
            "video_duration": value(row, "视频时长"),
            "content_review": content_review,
            "path": path,
            "media_dir": media_dir,
        }
        profiles.append(profile)

        media_folder = ROOT / media_dir
        media_folder.mkdir(parents=True, exist_ok=True)
        if process_media:
            for slot in range(1, 4):
                if not any((media_folder / f"{slot:02d}.{ext}").exists() for ext in IMAGE_EXTENSIONS):
                    media_warnings.append(f"{name}：缺少图片 {slot:02d}")
            if not (media_folder / "profile.mp4").exists():
                media_warnings.append(f"{name}：缺少视频 profile.mp4")

        if allow_index == "是":
            if city_by_slug[city_slug]["publish_status"] != "已发布" or city_by_slug[city_slug]["allow_index"] != "是":
                raise ValueError(f"资料表第 {row} 行所属城市尚未发布或不允许收录。")
            if content_review != "已完成":
                raise ValueError(f"资料表第 {row} 行允许收录前，内容审核状态必须改为“已完成”。")
            required_values = {
                "SEO标题": seo_title,
                "SEO描述": seo_description,
                "卡片/顶部简介": summary,
                "详情说明": profile["intro"],
                "正文1": profile["body_1"],
                "首次发布日期": profile["first_published"],
                "最后更新日期": profile["last_updated"],
            }
            missing = [label for label, text in required_values.items() if not text]
            if missing:
                raise ValueError(f"资料表第 {row} 行允许收录前必须填写：" + "、".join(missing))
            combined = " ".join(str(value) for value in profile.values())
            marker = next((item for item in TEMPLATE_MARKERS if item in combined), None)
            if marker:
                raise ValueError(f"资料表第 {row} 行仍包含模板文字“{marker}”，不能允许收录。")
            if profile["last_updated"] < profile["first_published"]:
                raise ValueError(f"资料表第 {row} 行最后更新日期不能早于首次发布日期。")
            if not any((media_folder / f"01.{ext}").exists() for ext in IMAGE_EXTENSIONS):
                raise ValueError(f"资料表第 {row} 行允许收录前必须添加真实主图 01。")
            if (media_folder / "profile.mp4").exists():
                if profile["video_duration"] and not re.fullmatch(r"PT(?=\d|\d.*[HMS])(?:\d+H)?(?:\d+M)?(?:\d+S)?", profile["video_duration"]):
                    raise ValueError(f"资料表第 {row} 行视频时长请使用 ISO 8601 格式，例如 PT1M30S。")

    if not profiles:
        raise ValueError("工作表中没有可更新的资料。")
    indexable = [profile for profile in profiles if profile["allow_index"] == "是"]
    for field, label in (("seo_title", "SEO标题"), ("seo_description", "SEO描述")):
        values: dict[str, str] = {}
        for profile in indexable:
            normalized = profile[field].strip().lower()
            if normalized in values:
                raise ValueError(f"允许收录的资料“{profile['name']}”与“{values[normalized]}”使用了重复{label}。")
            values[normalized] = profile["name"]

    profiles.sort(key=lambda item: (city_order[item["city_slug"]], item["display_order"], item["slug"]))
    if media_warnings:
        print("\n媒体检查提醒：")
        for warning in media_warnings:
            print("- " + warning)
    return profiles


def read_workbook(workbook_path: Path) -> tuple[list[dict], dict]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("缺少 openpyxl，请先运行：py -m pip install openpyxl") from exc

    if not SITE_CONTENT_PATH.exists():
        raise FileNotFoundError("缺少 site_content.json，无法确认整站字段。")
    existing = json.loads(SITE_CONTENT_PATH.read_text(encoding="utf-8"))
    formula_workbook = load_workbook(workbook_path, data_only=False)
    require_sheets(formula_workbook)
    if sync_profile_auto_formulas(formula_workbook):
        formula_workbook.save(workbook_path)

    workbook = load_workbook(workbook_path, data_only=True)
    require_sheets(workbook)

    home = read_settings_sheet(workbook, HOME_SHEET, set(existing["home"]))
    footer = read_settings_sheet(workbook, FOOTER_SHEET, set(existing["footer"]))
    cities = read_cities(workbook, existing["cities"])
    pages = read_fixed_pages(workbook, existing["pages"])
    profiles = read_profiles(workbook, cities, home.get("brand_name", "顶奢"))
    return profiles, {"home": home, "footer": footer, "cities": cities, "pages": pages}


def main() -> None:
    workbook_path = Path(sys.argv[1]).expanduser() if len(sys.argv) > 1 else DEFAULT_WORKBOOK
    if not workbook_path.is_absolute():
        workbook_path = ROOT / workbook_path
    if not workbook_path.exists():
        raise FileNotFoundError(f"找不到工作表：{workbook_path}")

    profiles, site_content = read_workbook(workbook_path)
    (ROOT / "profiles.json").write_text(
        json.dumps(profiles, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    SITE_CONTENT_PATH.write_text(
        json.dumps(site_content, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    subprocess.run([sys.executable, str(ROOT / "tools" / "build_site.py")], cwd=ROOT, check=True)
    subprocess.run([sys.executable, str(ROOT / "tools" / "check_site.py")], cwd=ROOT, check=True)
    print(f"\n整站更新完成：{len(profiles)} 个资料页、{len(site_content['cities'])} 个城市首页、5 个固定页面。")
    print(f"生成目录：{ROOT / 'dist'}")
    print("现在可以把完整项目上传到 GitHub。")


if __name__ == "__main__":
    try:
        main()
    except Exception as exc:
        print(f"\n更新失败：{exc}")
        raise SystemExit(1)
